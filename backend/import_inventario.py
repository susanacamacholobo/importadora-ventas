import openpyxl
import psycopg2
from dotenv import load_dotenv
import os

load_dotenv()

DATABASE_URL = os.getenv("DATABASE_URL")

# Parsear el connection string
import re
match = re.match(r'postgresql://(.+):(.+)@(.+):(\d+)/(.+)\?', DATABASE_URL)
if not match:
    match = re.match(r'postgresql://(.+):(.+)@(.+):(\d+)/(.+)', DATABASE_URL)

user, password, host, port, dbname = match.groups()

conn = psycopg2.connect(
    host=host, port=port, dbname=dbname,
    user=user, password=password, sslmode='require'
)
cur = conn.cursor()

print("Conectado a Supabase!")

# Cargar el Excel
wb = openpyxl.load_workbook(
    r"C:\Users\USUARIO\importadora-ventas\Inventario CONFIA a 24 de Agosto.xlsx",
    read_only=True,
    data_only=True
)

# Mapeo hoja → proveedor_id en la base de datos
PROVEEDORES_MAP = {
    'CASA IDEAS':    1,
    'AGENDAS':       2,
    'CHIKIPOOM COMPRA': 3,
    'IMBIMEX':       4,
    'MAKRO':         5,
    'MOXOS':         6,
    'KHOLBERG':      7,
    'DIEGO EID':     8,
    'EUROCASCOS':    9,
    'KIMPRO':        10,
    'ROXI':          11,
    'LISOFORT':      12,
    'GUABIRA':       13,
    'COICOM':        14,
    'AIDISA':        15,
    'UZ MUNDIAL':    16,
    'Confia':        17,
}

# IDs de ubicaciones
UBICACION = {
    'RECEPCION': 1,
    'ALMACEN':   2,
    'SHOWROOM':  3,
    'SEPARADO':  4,
}

total_insertados = 0
total_errores = 0

for sheet_name in wb.sheetnames:
    proveedor_id = PROVEEDORES_MAP.get(sheet_name)
    if not proveedor_id:
        print(f"  Saltando hoja: {sheet_name}")
        continue

    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    print(f"\nProcesando: {sheet_name} ({len(rows)} filas)")

    for row in rows[2:]:  # saltar cabeceras
        if not row or not row[0] or not row[1]:
            continue
        try:
            if len(row) < 8:               # ← LÍNEA NUEVA
                continue
            codigo   = str(row[0]).strip()
            nombre   = str(row[1]).strip()[:255]
            recep    = float(row[2]) if isinstance(row[2], (int, float)) else 0
            almacen  = float(row[3]) if isinstance(row[3], (int, float)) else 0
            showroom = float(row[4]) if isinstance(row[4], (int, float)) else 0
            separado = float(row[5]) if isinstance(row[5], (int, float)) else 0
            precio   = float(row[7]) if isinstance(row[7], (int, float)) else 0

            if not nombre or nombre == 'Nombre Producto':
                continue

            # Insertar producto
            cur.execute("""
                INSERT INTO productos (codigo_interno, nombre, precio_venta, precio_costo)
                VALUES (%s, %s, %s, %s)
                ON CONFLICT (codigo_interno) DO UPDATE
                SET nombre = EXCLUDED.nombre,
                    precio_venta = EXCLUDED.precio_venta
                RETURNING id
            """, (codigo, nombre, precio, precio * 0.7))

            prod_id = cur.fetchone()[0]

            # Relacionar con proveedor
            cur.execute("""
                INSERT INTO producto_proveedor (producto_id, proveedor_id, precio_proveedor)
                VALUES (%s, %s, %s)
                ON CONFLICT (producto_id, proveedor_id) DO NOTHING
            """, (prod_id, proveedor_id, precio))

            # Insertar stock por ubicación
            for ubicacion_id, cantidad in [
                (UBICACION['RECEPCION'], recep),
                (UBICACION['ALMACEN'],   almacen),
                (UBICACION['SHOWROOM'],  showroom),
                (UBICACION['SEPARADO'],  separado),
            ]:
                if cantidad > 0:
                    cur.execute("""
                        INSERT INTO stock (producto_id, ubicacion_id, cantidad, stock_minimo)
                        VALUES (%s, %s, %s, %s)
                        ON CONFLICT (producto_id, ubicacion_id) DO UPDATE
                        SET cantidad = EXCLUDED.cantidad
                    """, (prod_id, ubicacion_id, cantidad, 5))

            total_insertados += 1

        except Exception as e:
            print(f"  Error en fila {row[0]}: {type(e).__name__}: {e}")
            total_errores += 1
            conn.rollback()
            continue

        conn.commit()

print(f"\n✓ Importación completa!")
print(f"  Productos insertados: {total_insertados}")
print(f"  Errores: {total_errores}")

cur.close()
conn.close()